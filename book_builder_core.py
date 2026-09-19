# -*- coding: utf-8 -*-
"""
Full-Text Bestseller Book Generator
Title: Genspark(젠스파크) & 불편의 값: 누구나 쉽게 쓰는 올인원 AI 워크스페이스 60 & 430명 실증 비즈니스 생존 마스터북
"""

import sys
import os
import docx
from docx import Document
from docx.shared import Pt, Inches, RGBColor, Mm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn, nsdecls

sys.stdout.reconfigure(encoding='utf-8')

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=140, bottom=140, left=180, right=180):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def set_cell_borders(cell, top=None, bottom=None, left=None, right=None):
    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = parse_xml(f'<w:tcBorders {nsdecls("w")}/>')
    edges = {'top': top, 'bottom': bottom, 'left': left, 'right': right}
    for edge, val in edges.items():
        if val:
            tag = f'<w:{edge} {nsdecls("w")} w:val="{val.get("val", "single")}" w:sz="{val.get("sz", "4")}" w:space="0" w:color="{val.get("color", "auto")}"/>'
            tcBorders.append(parse_xml(tag))
        else:
            tag = f'<w:{edge} {nsdecls("w")} w:val="none"/>'
            tcBorders.append(parse_xml(tag))
    tcPr.append(tcBorders)

class BookBuilder:
    def __init__(self):
        self.doc = Document()
        self.setup_page()
        self.setup_styles()

    def setup_page(self):
        for section in self.doc.sections:
            section.page_width = Mm(152)
            section.page_height = Mm(225)
            section.top_margin = Mm(25)
            section.bottom_margin = Mm(25)
            section.left_margin = Mm(23)
            section.right_margin = Mm(21)
            section.different_first_page_header_footer = True

    def setup_styles(self):
        styles = self.doc.styles
        
        # Heading 1
        h1 = styles['Heading 1']
        h1.font.name = 'Malgun Gothic'
        h1.font.size = Pt(20)
        h1.font.bold = True
        h1.font.color.rgb = RGBColor(0x1E, 0x2A, 0x38)
        h1.paragraph_format.space_before = Pt(20)
        h1.paragraph_format.space_after = Pt(12)
        h1.paragraph_format.keep_with_next = True
        
        # Heading 2
        h2 = styles['Heading 2']
        h2.font.name = 'Malgun Gothic'
        h2.font.size = Pt(14)
        h2.font.bold = True
        h2.font.color.rgb = RGBColor(0x1E, 0x2A, 0x38)
        h2.paragraph_format.space_before = Pt(14)
        h2.paragraph_format.space_after = Pt(8)
        h2.paragraph_format.keep_with_next = True

        # Heading 3
        h3 = styles['Heading 3']
        h3.font.name = 'Malgun Gothic'
        h3.font.size = Pt(11.5)
        h3.font.bold = True
        h3.font.color.rgb = RGBColor(0xD9, 0x77, 0x06)
        h3.paragraph_format.space_before = Pt(10)
        h3.paragraph_format.space_after = Pt(6)
        h3.paragraph_format.keep_with_next = True

        # Heading 4
        h4 = styles['Heading 4']
        h4.font.name = 'Malgun Gothic'
        h4.font.size = Pt(11.0)
        h4.font.bold = True
        h4.font.color.rgb = RGBColor(0x47, 0x55, 0x69)
        h4.paragraph_format.space_before = Pt(8)
        h4.paragraph_format.space_after = Pt(4)
        h4.paragraph_format.keep_with_next = True

    def add_h1(self, text):
        p = self.doc.add_paragraph(text, style='Heading 1')
        p.paragraph_format.keep_with_next = True
        return p

    def add_h2(self, text):
        p = self.doc.add_paragraph(text, style='Heading 2')
        p.paragraph_format.keep_with_next = True
        return p

    def add_h3(self, text):
        p = self.doc.add_paragraph(text, style='Heading 3')
        p.paragraph_format.keep_with_next = True
        return p

    def add_h4(self, text):
        p = self.doc.add_paragraph(text, style='Heading 4')
        p.paragraph_format.keep_with_next = True
        return p

    def add_body(self, text, bold_prefix=None, indent=True):
        p = self.doc.add_paragraph()
        p.paragraph_format.line_spacing = 1.75
        p.paragraph_format.space_after = Pt(8.5)
        if indent:
            p.paragraph_format.first_line_indent = Pt(10.5)
        else:
            p.paragraph_format.first_line_indent = Pt(0)
            
        if bold_prefix:
            r_bold = p.add_run(bold_prefix)
            r_bold.font.name = 'Batang'
            r_bold.font.size = Pt(10.5)
            r_bold.font.bold = True
            r_bold.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A) # Slate Navy
            
        r_text = p.add_run(text)
        r_text.font.name = 'Batang'
        r_text.font.size = Pt(10.5)
        r_text.font.color.rgb = RGBColor(0x2A, 0x34, 0x39) # Warm Charcoal
        return p

    def add_callout(self, title, content_lines, theme='gold'):
        """
        theme: 'gold', 'navy', 'slate'
        """
        colors = {
            'gold': {'border': 'D97706', 'bg': 'FFFBEB', 'title': RGBColor(0xB4, 0x53, 0x09)},
            'navy': {'border': '1E2A38', 'bg': 'F8FAFC', 'title': RGBColor(0x1E, 0x2A, 0x38)},
            'slate': {'border': '475569', 'bg': 'F1F5F9', 'title': RGBColor(0x33, 0x41, 0x55)}
        }
        cfg = colors.get(theme, colors['gold'])
        
        tbl = self.doc.add_table(rows=1, cols=1)
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        cell = tbl.cell(0, 0)
        set_cell_background(cell, cfg['bg'])
        set_cell_margins(cell, top=140, bottom=140, left=180, right=180)
        set_cell_borders(cell, left={'val': 'single', 'sz': '24', 'color': cfg['border']},
                              top={'val': 'single', 'sz': '4', 'color': 'E2E8F0'},
                              bottom={'val': 'single', 'sz': '4', 'color': 'E2E8F0'},
                              right={'val': 'single', 'sz': '4', 'color': 'E2E8F0'})
        
        p = cell.paragraphs[0]
        p.paragraph_format.line_spacing = 1.3
        p.paragraph_format.space_after = Pt(4)
        r_title = p.add_run(f"■ {title}\n")
        r_title.font.name = 'Malgun Gothic'
        r_title.font.bold = True
        r_title.font.size = Pt(10.5)
        r_title.font.color.rgb = cfg['title']
        
        for idx, line in enumerate(content_lines):
            p2 = cell.add_paragraph()
            p2.paragraph_format.line_spacing = 1.4
            p2.paragraph_format.space_after = Pt(3 if idx < len(content_lines)-1 else 0)
            r_line = p2.add_run(line)
            r_line.font.name = 'Batang'
            r_line.font.size = Pt(9.5)
            r_line.font.color.rgb = RGBColor(0x33, 0x41, 0x55)
            
        p_space = self.doc.add_paragraph()
        p_space.paragraph_format.space_after = Pt(6)

    def add_table_data(self, headers, rows, col_widths=None):
        tbl = self.doc.add_table(rows=len(rows)+1, cols=len(headers))
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        
        # Header row
        hdr_cells = tbl.rows[0].cells
        for i, title in enumerate(headers):
            hdr_cells[i].text = title
            set_cell_background(hdr_cells[i], '1E2A38')
            set_cell_margins(hdr_cells[i], top=120, bottom=120, left=140, right=140)
            p = hdr_cells[i].paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.runs[0].font.name = 'Malgun Gothic'
            p.runs[0].font.bold = True
            p.runs[0].font.size = Pt(9.5)
            p.runs[0].font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
            
        # Data rows
        for r_idx, r_data in enumerate(rows):
            row_cells = tbl.rows[r_idx+1].cells
            bg_color = 'F8FAFC' if r_idx % 2 == 1 else 'FFFFFF'
            for c_idx, val in enumerate(r_data):
                row_cells[c_idx].text = str(val)
                set_cell_background(row_cells[c_idx], bg_color)
                set_cell_margins(row_cells[c_idx], top=100, bottom=100, left=120, right=120)
                set_cell_borders(row_cells[c_idx], 
                                 top={'val': 'single', 'sz': '4', 'color': 'CBD5E1'},
                                 bottom={'val': 'single', 'sz': '4', 'color': 'CBD5E1'},
                                 left={'val': 'single', 'sz': '4', 'color': 'E2E8F0'},
                                 right={'val': 'single', 'sz': '4', 'color': 'E2E8F0'})
                p = row_cells[c_idx].paragraphs[0]
                p.paragraph_format.line_spacing = 1.3
                p.paragraph_format.space_after = Pt(0)
                if len(p.runs) > 0:
                    p.runs[0].font.name = 'Batang'
                    p.runs[0].font.size = Pt(9.0)
                    p.runs[0].font.color.rgb = RGBColor(0x33, 0x41, 0x55)
                    
        p_space = self.doc.add_paragraph()
        p_space.paragraph_format.space_after = Pt(8)

    def add_page_break(self):
        self.doc.add_page_break()

    def save(self, output_path):
        self.doc.save(output_path)
        print(f"Master Book successfully saved to: {output_path}")

print("BookBuilder core framework ready.")
